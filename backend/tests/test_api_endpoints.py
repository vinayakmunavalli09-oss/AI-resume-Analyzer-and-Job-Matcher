"""
Integration tests for FastAPI endpoints using TestClient.
"""
import io
import unittest
from fastapi.testclient import TestClient
import docx

import main


def make_test_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Deepak Patel - Fresher Resume", 0)
    doc.add_paragraph("Email: deepak@example.com | Phone: 9876543210 | Location: Bengaluru, Karnataka")
    doc.add_paragraph("Education: B.Tech in Information Science, 2026")
    doc.add_paragraph("Skills: Python, FastAPI, SQL, React, Git, Docker, Problem Solving")
    doc.add_paragraph("Academic Projects: E-Commerce API built using FastAPI, PostgreSQL, and Python")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestApiEndpoints(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(main.app)

    def test_health_check(self):
        resp = self.client.get("/health")
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(data["status"], "ok")
        self.assertIn("nvidia_configured", data)
        self.assertIn("storage_mode", data)

    def test_list_jobs_and_filters(self):
        # 1. Fetch all
        resp = self.client.get("/api/v1/jobs?page_size=60")
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertGreaterEqual(data["total"], 20)
        self.assertTrue(len(data["jobs"]) >= 20)

        # 2. Filter by location
        resp_loc = self.client.get("/api/v1/jobs?location=Bengaluru")
        self.assertEqual(resp_loc.status_code, 200)
        for job in resp_loc.json()["jobs"]:
            self.assertTrue(
                "bengaluru" in job["location"].lower() or "bangalore" in job["location"].lower()
            )

        # 3. Filter by role category
        resp_cat = self.client.get("/api/v1/jobs?role_category=Digital%20Marketing")
        self.assertEqual(resp_cat.status_code, 200)
        for job in resp_cat.json()["jobs"]:
            self.assertIn("marketing", (job["role_category"] + " " + job["title"]).lower())

    def test_get_single_job(self):
        # Fetch first job ID
        list_resp = self.client.get("/api/v1/jobs?page_size=1")
        job_id = list_resp.json()["jobs"][0]["id"]

        resp = self.client.get(f"/api/v1/jobs/{job_id}")
        self.assertEqual(resp.status_code, 200)
        job = resp.json()
        self.assertEqual(job["id"], job_id)
        self.assertTrue(job["is_demo"])
        self.assertEqual(job["source_label"], "Demo Data")

    def test_resume_analysis_and_matching_workflow(self):
        docx_bytes = make_test_docx()
        files = {
            "file": (
                "deepak_resume.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        session_id = "test-private-session-capability-0f5d5d3f6fe84efc"
        headers = {"X-Session-Id": session_id}
        files["session_id"] = (None, session_id)

        # 1. Analyze resume
        resp = self.client.post("/api/v1/resumes/analyze", files=files, headers=headers)
        self.assertEqual(resp.status_code, 200)
        analysis = resp.json()
        analysis_id = analysis["analysis_id"]
        self.assertTrue(bool(analysis_id))
        self.assertIn("Python", analysis["structured_profile"]["skills"])
        self.assertGreater(analysis["feedback"]["overall_score"], 0)

        # 2. Retrieve analysis by ID
        get_resp = self.client.get(f"/api/v1/resumes/{analysis_id}", headers=headers)
        self.assertEqual(get_resp.status_code, 200)
        self.assertEqual(get_resp.json()["analysis_id"], analysis_id)

        # 3. Candidate Review and Correction (PUT)
        updated_profile = dict(analysis["structured_profile"])
        updated_profile["skills"].append("Kubernetes")
        put_resp = self.client.put(f"/api/v1/resumes/{analysis_id}", json=updated_profile, headers=headers)
        self.assertEqual(put_resp.status_code, 200)
        self.assertIn("Kubernetes", put_resp.json()["structured_profile"]["skills"])

        # 4. Calculate matches
        match_req = {"analysis_id": analysis_id}
        match_resp = self.client.post("/api/v1/matches", json=match_req, headers=headers)
        self.assertEqual(match_resp.status_code, 200)
        matches_data = match_resp.json()
        self.assertGreater(matches_data["total"], 0)
        first_match = matches_data["matches"][0]
        self.assertIn("overall_score", first_match)
        self.assertIn("skills", first_match["score_breakdown"])

        # 5. Get skill gaps for top match
        match_id = first_match["match_id"]
        gap_resp = self.client.get(f"/api/v1/matches/{match_id}/skill-gaps", headers=headers)
        self.assertEqual(gap_resp.status_code, 200)
        gap_data = gap_resp.json()
        self.assertIn("matched_skills", gap_data)
        self.assertIn("missing_required", gap_data)

        # 6. Privacy Deletion
        del_resp = self.client.delete(f"/api/v1/resumes/{analysis_id}", headers=headers)
        self.assertEqual(del_resp.status_code, 200)
        self.assertTrue(del_resp.json()["success"])

        # Verify 404 after deletion
        get_deleted = self.client.get(f"/api/v1/resumes/{analysis_id}", headers=headers)
        self.assertEqual(get_deleted.status_code, 404)

    def test_resume_analysis_requires_owning_session(self):
        docx_bytes = make_test_docx()
        files = {
            "file": (
                "resume.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
            "session_id": (None, "session-owner-capability-1234567890"),
        }
        owner_headers = {"X-Session-Id": "session-owner-capability-1234567890"}
        response = self.client.post("/api/v1/resumes/analyze", files=files, headers=owner_headers)
        self.assertEqual(response.status_code, 200)
        analysis_id = response.json()["analysis_id"]
        stranger = self.client.get(
            f"/api/v1/resumes/{analysis_id}",
            headers={"X-Session-Id": "different-session-capability-1234567890"},
        )
        self.assertEqual(stranger.status_code, 404)

    def test_saved_jobs_lifecycle(self):
        # 1. Bookmark a job
        save_resp = self.client.post(
            "/api/v1/saved-jobs",
            json={"job_id": "job-001", "notes": "Applied on portal", "status": "applied"},
            headers={"X-User-Id": "test_user_99"},
        )
        self.assertEqual(save_resp.status_code, 201)

        # 2. List saved jobs
        list_resp = self.client.get(
            "/api/v1/saved-jobs",
            headers={"X-User-Id": "test_user_99"},
        )
        self.assertEqual(list_resp.status_code, 200)
        jobs = list_resp.json()
        self.assertGreaterEqual(len(jobs), 1)
        self.assertEqual(jobs[0]["job_id"], "job-001")
        self.assertEqual(jobs[0]["application_status"], "applied")

        # 3. Update status (e.g. to interview)
        patch_resp = self.client.patch(
            "/api/v1/saved-jobs/job-001",
            json={"status": "interview"},
            headers={"X-User-Id": "test_user_99"},
        )
        self.assertEqual(patch_resp.status_code, 200)
        self.assertEqual(patch_resp.json()["application_status"], "interview")

        # 4. Remove saved job
        del_resp = self.client.delete(
            "/api/v1/saved-jobs/job-001",
            headers={"X-User-Id": "test_user_99"},
        )
        self.assertEqual(del_resp.status_code, 200)

    def test_profile_preferences(self):
        profile_data = {
            "name": "Priya Nair",
            "email": "priya@example.com",
            "location": "Pune, Maharashtra",
            "target_roles": ["Business Analyst Trainee", "Data Associate"],
            "preferred_locations": ["Pune", "Mumbai", "Remote"],
            "skills": ["SQL", "Power BI", "Excel"],
        }
        put_resp = self.client.put(
            "/api/v1/profile",
            json=profile_data,
            headers={"X-User-Id": "priya_test"},
        )
        self.assertEqual(put_resp.status_code, 200)
        saved = put_resp.json()
        self.assertEqual(saved["name"], "Priya Nair")

        get_resp = self.client.get(
            "/api/v1/profile",
            headers={"X-User-Id": "priya_test"},
        )
        self.assertEqual(get_resp.status_code, 200)
        self.assertEqual(get_resp.json()["location"], "Pune, Maharashtra")


if __name__ == "__main__":
    unittest.main()
