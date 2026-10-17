import io
import httpx
import docx

BASE = "http://127.0.0.1:8000"

def run_checks():
    # 1. Health
    h = httpx.get(f"{BASE}/health").json()
    print("HEALTH:", h)
    assert h["status"] == "ok"

    # 2. Jobs
    jobs_res = httpx.get(f"{BASE}/api/v1/jobs?page=1&page_size=5").json()
    print("JOBS TOTAL:", jobs_res["total"], "RETURNED:", len(jobs_res["jobs"]))
    assert jobs_res["total"] >= 20

    # 3. Create resume docx in memory
    doc = docx.Document()
    doc.add_heading("Deepak Patel - Fresher Resume", 0)
    doc.add_paragraph("Email: deepak@example.com | Phone: 9876543210 | Location: Bengaluru, Karnataka")
    doc.add_paragraph("Education: B.Tech in Information Science, 2026")
    doc.add_paragraph("Skills: Python, FastAPI, SQL, React, Git, Docker, Problem Solving")
    doc.add_paragraph("Academic Projects: E-Commerce API built using FastAPI, PostgreSQL, and Python")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    # 4. Upload & Analyze Resume
    session_id = "test-session-live-verify-1234567890abcdef"
    files = {
        "file": (
            "deepak_resume.docx",
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    headers = {"X-Session-Id": session_id}
    data = {"session_id": session_id}
    res_analyze = httpx.post(f"{BASE}/api/v1/resumes/analyze", files=files, data=data, headers=headers, timeout=120.0)
    print("RESUME ANALYZE STATUS:", res_analyze.status_code)
    assert res_analyze.status_code == 200, res_analyze.text
    analysis = res_analyze.json()
    analysis_id = analysis["analysis_id"]
    print("ANALYSIS ID:", analysis_id)
    print("EXTRACTED SKILLS:", analysis["structured_profile"]["skills"])
    print("OVERALL FEEDBACK SCORE:", analysis["feedback"]["overall_score"])

    # 5. Matching
    match_res = httpx.post(f"{BASE}/api/v1/matches", json={"analysis_id": analysis_id}, headers=headers).json()
    print("MATCHES FOUND:", match_res["total"])
    assert match_res["total"] > 0
    top_match = match_res["matches"][0]
    match_id = top_match["match_id"]
    print("TOP MATCH JOB ID:", top_match["job_id"], "SCORE:", top_match["overall_score"], "EXPLANATION:", top_match["explanation"][:80])

    # 6. Skill Gap
    gap_res = httpx.get(f"{BASE}/api/v1/matches/{match_id}/skill-gaps", headers=headers).json()
    print("SKILL GAP COVERAGE %:", gap_res["coverage_percent"])
    print("MATCHED SKILLS:", gap_res["matched_skills"])
    print("MISSING REQUIRED:", [item["skill"] for item in gap_res["missing_required"]])

    # 7. Saved jobs
    save_res = httpx.post(
        f"{BASE}/api/v1/saved-jobs",
        json={"job_id": top_match["job_id"], "status": "applied"},
        headers={"X-User-Id": "live-user"},
    ).json()
    print("SAVED JOB STATUS:", save_res["application_status"])

    # 8. Clean up resume
    del_res = httpx.delete(f"{BASE}/api/v1/resumes/{analysis_id}", headers=headers).json()
    print("DELETE RESUME SUCCESS:", del_res["success"])

    print("ALL LIVE SERVER VERIFICATIONS PASSED!")

if __name__ == "__main__":
    run_checks()
