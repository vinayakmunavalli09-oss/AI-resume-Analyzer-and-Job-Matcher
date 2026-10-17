"""
Resume file parser: extracts plain text from PDF and DOCX uploads.

Security notes:
- Only PDF and DOCX are accepted (no executable formats).
- Files are read entirely in memory; no disk writes.
- We never execute the uploaded file content.
"""
import io
import os
import logging
from pathlib import Path

from fastapi import UploadFile

from app.models.resume import ParseResult
from app.core.exceptions import (
    ResumeParseError,
    FileTooLargeError,
    UnsupportedFileTypeError,
)
from app.core.security import ALLOWED_MIME_TYPES, ALLOWED_EXTENSIONS
from config import get_settings

logger = logging.getLogger(__name__)


async def parse_resume(file: UploadFile) -> ParseResult:
    """
    Validate and parse an uploaded resume file.

    Args:
        file: FastAPI UploadFile (PDF or DOCX).

    Returns:
        ParseResult with extracted_text, word_count, and any warnings.

    Raises:
        FileTooLargeError: if file exceeds max_file_size_mb.
        UnsupportedFileTypeError: if file is not PDF/DOCX.
        ResumeParseError: if extraction fails.
    """
    settings = get_settings()
    max_bytes = settings.max_file_size_mb * 1024 * 1024
    warnings: list[str] = []

    # ── Read file into memory ────────────────────────────────────────────────
    raw_bytes = await file.read()
    if not raw_bytes:
        raise ResumeParseError("The uploaded file is empty.", status_code=400)
    if len(raw_bytes) > max_bytes:
        raise FileTooLargeError(settings.max_file_size_mb)

    # ── Validate file type ───────────────────────────────────────────────────
    content_type = (file.content_type or "").lower().split(";")[0].strip()
    filename = file.filename or ""
    suffix = Path(filename).suffix.lower()

    is_valid_mime = content_type in ALLOWED_MIME_TYPES
    is_valid_ext = suffix in ALLOWED_EXTENSIONS

    if not (is_valid_mime or is_valid_ext):
        raise UnsupportedFileTypeError(content_type or suffix)

    # Determine actual format from content (magic bytes) as secondary guard
    if raw_bytes[:4] == b"%PDF":
        actual_format = "pdf"
    elif raw_bytes[:2] == b"PK":  # ZIP-based DOCX
        actual_format = "docx"
    else:
        raise UnsupportedFileTypeError(content_type)

    # ── Extract text ─────────────────────────────────────────────────────────
    if actual_format == "pdf":
        extracted_text, pdf_warnings = _extract_pdf(raw_bytes)
        warnings.extend(pdf_warnings)
    else:
        extracted_text, docx_warnings = _extract_docx(raw_bytes)
        warnings.extend(docx_warnings)

    if not extracted_text.strip():
        raise ResumeParseError(
            "No readable text could be extracted from the file. "
            "If this is a scanned PDF, please provide a text-based version.",
            status_code=400,
        )

    word_count = len(extracted_text.split())
    if word_count < 20:
        warnings.append(
            f"Very little text extracted ({word_count} words). "
            "The resume may be mostly images or have unusual formatting."
        )

    logger.info(
        "Parsed resume '%s': %d words, format=%s, warnings=%d",
        filename,
        word_count,
        actual_format,
        len(warnings),
    )
    return ParseResult(
        extracted_text=extracted_text,
        word_count=word_count,
        warnings=warnings,
    )


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _extract_pdf(raw_bytes: bytes) -> tuple[str, list[str]]:
    """Extract text from a PDF byte string using pypdf."""
    warnings: list[str] = []
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw_bytes))
        if reader.is_encrypted:
            warnings.append("PDF is encrypted/password-protected; extraction may be incomplete.")
            try:
                reader.decrypt("")
            except Exception:
                raise ResumeParseError(
                    "The PDF is password-protected and cannot be read.", status_code=400
                )

        pages_text: list[str] = []
        for page in reader.pages:
            try:
                text = page.extract_text() or ""
                pages_text.append(text)
            except Exception as exc:
                logger.warning("Could not extract one PDF page")
                warnings.append("One page could not be read; its text may be incomplete.")

        full_text = "\n".join(pages_text)
        return full_text, warnings

    except ResumeParseError:
        raise
    except Exception as exc:
        logger.warning("PDF extraction failed")
        raise ResumeParseError("Failed to read the PDF. Please check that it is a valid, unencrypted document.", status_code=400) from exc


def _extract_docx(raw_bytes: bytes) -> tuple[str, list[str]]:
    """Extract text from a DOCX byte string using python-docx."""
    warnings: list[str] = []
    try:
        import docx

        doc = docx.Document(io.BytesIO(raw_bytes))
        paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]

        # Also capture text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())

        full_text = "\n".join(paragraphs)
        return full_text, warnings

    except Exception as exc:
        logger.warning("DOCX extraction failed")
        raise ResumeParseError("Failed to read the DOCX. Please check that it is a valid document.", status_code=400) from exc
